"""Read existing delivered Excel rows including unpriced descriptions; do not edit Excel."""
import json, hashlib, sys
from pathlib import Path
from openpyxl import load_workbook

manifest=json.loads(Path(sys.argv[1]).read_text())
out=[]
for entry in manifest:
    path=Path(entry['folder'])/entry['rabFile']
    wb=load_workbook(path,read_only=True,data_only=True)
    ws=wb[entry['sheet']]
    rows=[]
    for index, cells in enumerate(ws.iter_rows(min_row=3,max_row=entry['totalRow']-1,max_col=21,values_only=True),3):
        description=str(cells[1] or '').strip()
        if not description:
            continue
        item=index in entry['itemRows']
        if not item:
            assert not cells[6] and all(cells[c-1] in (None,'',0) for c in [4,5,6,8,9,10,11,12,13,14,19,20,21]), f'Ambiguous unpriced row: {path}:{index}'
        rows.append({'source_row':index,'urutan':len(rows)+1,'nomor':str(cells[0] or ''),'uraian':description,'type':'ITEM' if item else 'KETERANGAN','satuan':str(cells[6] or ''),
                     'jumlah_hps':cells[18] if item else None,'jumlah_penawaran':cells[19] if item else None,'jumlah_negoisasi':cells[20] if item else None})
    assert sum(r['type']=='ITEM' for r in rows)==len(entry['itemRows'])
    out.append({'id':entry['id'],'path':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'sheet':entry['sheet'],'rows':rows})
    wb.close()
Path(sys.argv[2]).write_text(json.dumps({'contracts':out},ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'packages':len(out),'items':sum(r['type']=='ITEM' for c in out for r in c['rows']),'notes':sum(r['type']=='KETERANGAN' for c in out for r in c['rows'])}))
