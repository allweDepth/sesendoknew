#!/usr/bin/env python3
"""Ekstrak kontrak, rekanan, dan RAB dari workbook Numbers/Excel sebagai data."""
from __future__ import annotations
import json,re,sys
from datetime import date,datetime,timedelta
from pathlib import Path
from openpyxl import load_workbook

def val(v):
    if isinstance(v,(datetime,date)): return v.isoformat()[:10]
    return v
def text(v): return re.sub(r"\s+"," ",str(v or "")).strip()
def num(v):
    try:return float(v or 0)
    except:return 0.0
def rows(path): return list(load_workbook(path,read_only=True,data_only=True).active.iter_rows(values_only=True))

contract_rows=rows(sys.argv[1]); headers=[text(x) for x in contract_rows[3]]
contracts=[]
for excel_row,row in enumerate(contract_rows[4:],5):
    if not row[1]:continue
    r=list(row)+[None]*80
    contracts.append({"excel_row":excel_row,"row_number":text(r[1]),"package":text(r[2]),"sub_activity_name":text(r[3]),"location_package":text(r[4]),"budget":num(r[9]),"hps":num(r[10]),"offer":num(r[11]),"negotiated":num(r[12]),"vendor_label":text(r[13]),"vendor_excel_id":int(num(r[58])) if r[58] is not None else None,"director":text(r[14]),"position":text(r[15]),"npwp":text(r[16]),"address":text(r[17]),"notary_number":text(r[18]),"notary_date":val(r[19]),"notary_name":text(r[20]),"bank_account":text(r[21]),"bank_account_name":text(r[22]),"duration_days":int(num(r[23])),"funding":text(r[24]),"contract_value":num(r[25]),"offer_number":text(r[26]),"offer_date":val(r[27]),"maintenance_days":int(num(r[28])),"contract_type":text(r[29]),"category":text(r[30]),"sub_activity_code":text(r[31]),"invitation_date":val(r[32]),"opening_date":val(r[33]),"evaluation_date":val(r[34]),"negotiation_date":val(r[35]),"spbbj_date":val(r[36]),"contract_date":val(r[37]),"spmk_date":val(r[38]),"bahp_number":text(r[39]),"bahp_date":val(r[40]),"determination_number":text(r[41]),"determination_date":val(r[42]),"announcement_number":text(r[43]),"announcement_date":val(r[44]),"negotiation_number":text(r[45]),"negotiation_proof_date":val(r[46]),"evaluation_value":num(r[47]),"contract_number":text(r[48]),"spmk_number":text(r[49]),"addendum_number":text(r[52]),"addendum_date":val(r[53]),"email":text(r[56]),"procurement_method":text(r[60]),"submission_method":text(r[61]),"qualification_method":text(r[62]),"evaluation_method":text(r[63]),"planning_consultant":text(r[66]),"supervision_consultant":text(r[67]),"performance_guarantee":num(r[70]),"technical_team":text(r[71]),"notes":text(r[72])})

vendor_rows=rows(sys.argv[2]);vendors=[]
for row in vendor_rows[2:]:
    r=list(row)+[None]*30
    if not isinstance(r[0],(int,float)) or not text(r[1]):continue
    vendors.append({"excel_id":int(r[0]),"name":text(r[1]),"status":text(r[2]),"address":text(r[3]),"email":text(r[4]),"director":text(r[5]),"position":text(r[6]),"id_card":text(r[7]),"phone":text(r[8]),"npwp":text(r[9]),"deed_number":text(r[10]),"deed_date":val(r[11]),"deed_notary":text(r[12]),"deed_place":text(r[13]),"change_deed_number":text(r[14]),"change_deed_date":val(r[15]),"change_deed_notary":text(r[16]),"change_deed_place":text(r[17]),"bank_account":text(r[18]),"bank_account_name":text(r[19]),"notes":text(r[23]),"proxy":text(r[24]),"sort":int(num(r[25])) if r[25] is not None else None,"description":text(r[26])})

rab_book=load_workbook(sys.argv[3],read_only=True,data_only=True); rabs={}; rab_candidates={}
contract_numbers={c["row_number"] for c in contracts}
for sheet in rab_book.worksheets:
    match=next((n for n in contract_numbers if re.match(rf"^{re.escape(n)}(?:\D|$)",sheet.title)),None)
    if not match:continue
    detail=[]
    for excel_row,row in enumerate(sheet.iter_rows(values_only=True),1):
        r=list(row)+[None]*14;description=text(r[1])
        if excel_row<=2 or not description or not isinstance(r[0],(int,float)) or num(r[0])<=0:continue
        tax=num(r[10]);qh,qo,qn=num(r[3]),num(r[4]),num(r[5]);ph,po,pn=num(r[7]),num(r[8]),num(r[9])
        detail.append({"excel_row":excel_row,"number":text(r[0]),"description":description,"quantity_hps":qh,"quantity_offer":qo,"quantity_negotiated":qn,"unit":text(r[6]),"unit_price_hps":ph,"unit_price_offer":po,"unit_price_negotiated":pn,"tax":tax,"amount_hps":qh*ph*(1+tax/100),"amount_offer":qo*po*(1+tax/100),"amount_negotiated":qn*pn*(1+tax/100)})
    score=sum(1 for x in detail if x["amount_negotiated"]>0)
    if score>(rab_candidates.get(match,(0,None))[0]):rab_candidates[match]=(score,{"sheet":sheet.title,"items":detail})
rabs={key:value for key,(_,value) in rab_candidates.items() if value}

payload={"scope":{"kd_wilayah":"76.01","kd_opd":"1.03.0.00.0.00.01.0000","year":2026},"contracts":contracts,"vendors":vendors,"rabs":rabs}
Path(sys.argv[4]).write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8")
print(json.dumps({"contracts":len(contracts),"vendors":len(vendors),"rab_contracts":len(rabs),"rab_items":sum(len(x['items']) for x in rabs.values()),"missing_rab":[x['row_number'] for x in contracts if x['row_number'] not in rabs]},ensure_ascii=False))
